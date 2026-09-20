"""共享导出服务接口测试：A-06、I-03 精神与依赖缺失降级。"""

from __future__ import annotations

import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.core.config import settings
from app.db.models import ExportRecord
from app.services.export import libreoffice

TEMPLATE_ID = "classic_single_column"


def docx_text(content: bytes) -> str:
    """提取 DOCX 文本；制表位拼接的列统一还原成 `` | `` 便于断言。"""

    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text.replace("\t", " | ") for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.extend(cell.text.replace("\t", " | ") for cell in row.cells)
    return "\n".join(paragraphs)


def test_export_templates_endpoint(client: TestClient) -> None:
    response = client.get("/api/export-templates")

    assert response.status_code == 200
    items = response.json()["data"]["items"]
    assert [item["id"] for item in items] == [TEMPLATE_ID]
    assert items[0]["formats"] == ["docx", "pdf"]
    assert "微软雅黑" in items[0]["description"] or "参考简历" in items[0]["description"]


def test_a06_export_docx_contains_version_content(client: TestClient, resume_factory) -> None:
    _, version_id = resume_factory()

    response = client.post(
        f"/api/resume-versions/{version_id}/exports",
        json={"templateId": TEMPLATE_ID, "format": "docx"},
    )

    assert response.status_code == 201, response.text
    data = response.json()["data"]
    assert data["status"] == "completed"
    assert data["templateId"] == TEMPLATE_ID
    assert data["downloadUrl"] == f"/api/exports/{data['id']}/download"

    download = client.get(data["downloadUrl"])
    assert download.status_code == 200
    assert download.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )

    text = docx_text(download.content)
    # 基本信息按「标签：值」输出
    assert "姓名：李小明" in text
    assert "手机：13800000000" in text
    assert "邮箱：lixiaoming@example.com" in text
    assert "求职意向：后端开发工程师" in text

    # 模块顺序与版本一致
    assert text.index("教育背景") < text.index("实习经历") < text.index("项目经历") < text.index("技能与工具")

    # 条目为「名称 | 角色 | 起止时间」三列
    assert "浙江大学 | 计算机科学与技术 本科 | 2019.09 - 2023.06" in text
    assert "杭州云帆科技有限公司 | 后端开发实习生 | 2024.07 - 2024.12" in text
    assert "ResumeMatch | 后端开发 | 2026.09 - 2026.12" in text

    # 描述与技能行完整保留
    assert "参与订单系统重构，接口平均响应时间下降 40%；" in text
    assert "项目背景：面向求职场景的简历匹配系统。" in text
    assert "编程：Python，Java，Git" in text
    assert "AI工具：Trae，VSCode" in text

    # 占位符不得残留
    assert "{{" not in text and "}}" not in text


def test_export_follows_module_and_item_order(client: TestClient, resume_factory) -> None:
    document = {
        "basicInfo": {"name": "顺序校验"},
        "sections": [
            {
                "id": "section_project",
                "type": "project",
                "title": "项目经历",
                "order": 0,
                "items": [
                    {"id": "item_project_1", "order": 0, "fields": {"name": "第一"}, "content": ""},
                    {"id": "item_project_2", "order": 1, "fields": {"name": "第二"}, "content": ""},
                ],
            },
            {
                "id": "section_education",
                "type": "education",
                "title": "教育背景",
                "order": 1,
                "items": [
                    {
                        "id": "item_education_1",
                        "order": 0,
                        "fields": {"school": "某大学", "startDate": "2019-09", "endDate": "2023-06"},
                        "content": "",
                    }
                ],
            },
        ],
    }
    _, version_id = resume_factory(document)

    created = client.post(
        f"/api/resume-versions/{version_id}/exports",
        json={"templateId": TEMPLATE_ID, "format": "docx"},
    ).json()["data"]
    text = docx_text(client.get(created["downloadUrl"]).content)

    assert text.index("项目经历") < text.index("教育背景")
    assert text.index("第一") < text.index("第二")


def test_i03_same_version_template_format_is_consistent(client: TestClient, resume_factory) -> None:
    _, version_id = resume_factory()
    body = {"templateId": TEMPLATE_ID, "format": "docx"}

    first = client.post(f"/api/resume-versions/{version_id}/exports", json=body).json()["data"]
    second = client.post(f"/api/resume-versions/{version_id}/exports", json=body).json()["data"]

    assert first["id"] != second["id"]
    first_text = docx_text(client.get(first["downloadUrl"]).content)
    second_text = docx_text(client.get(second["downloadUrl"]).content)
    assert first_text == second_text


def test_version_snapshot_and_export_stay_in_sync(client: TestClient, resume_factory) -> None:
    """导出结果来自不可变版本：改草稿不影响再次导出的内容。"""

    resume_id, version_id = resume_factory()
    body = {"templateId": TEMPLATE_ID, "format": "docx"}
    first = client.post(f"/api/resume-versions/{version_id}/exports", json=body).json()["data"]

    client.put(
        f"/api/resumes/{resume_id}/draft",
        json={
            "document": {
                "basicInfo": {"name": "草稿已改"},
                "sections": [],
            }
        },
    )
    second = client.post(f"/api/resume-versions/{version_id}/exports", json=body).json()["data"]

    assert "姓名：李小明" in docx_text(client.get(first["downloadUrl"]).content)
    assert "姓名：李小明" in docx_text(client.get(second["downloadUrl"]).content)
    assert "草稿已改" not in docx_text(client.get(second["downloadUrl"]).content)


def test_export_records_are_listed_by_status(client: TestClient, resume_factory) -> None:
    _, version_id = resume_factory()
    created = client.post(
        f"/api/resume-versions/{version_id}/exports",
        json={"templateId": TEMPLATE_ID, "format": "docx"},
    ).json()["data"]

    response = client.get(f"/api/exports/{created['id']}")

    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["status"] == "completed"
    assert payload["resumeVersionId"] == version_id
    assert payload["completedAt"] is not None


def test_export_unknown_version(client: TestClient) -> None:
    response = client.post(
        "/api/resume-versions/version_missing/exports",
        json={"templateId": TEMPLATE_ID, "format": "docx"},
    )

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "VERSION_NOT_FOUND"


def test_export_unknown_template(client: TestClient, resume_factory) -> None:
    _, version_id = resume_factory()

    response = client.post(
        f"/api/resume-versions/{version_id}/exports",
        json={"templateId": "not_exists", "format": "docx"},
    )

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "TEMPLATE_NOT_FOUND"


def test_export_invalid_format(client: TestClient, resume_factory) -> None:
    _, version_id = resume_factory()

    response = client.post(
        f"/api/resume-versions/{version_id}/exports",
        json={"templateId": TEMPLATE_ID, "format": "rtf"},
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "VALIDATION_ERROR"


def test_unknown_export_download(client: TestClient) -> None:
    response = client.get("/api/exports/export_missing")

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "EXPORT_NOT_FOUND"


def test_pdf_without_libreoffice_fails_but_is_recorded(
    client: TestClient, session_factory, resume_factory, monkeypatch
) -> None:
    _, version_id = resume_factory()
    monkeypatch.setattr(settings, "libreoffice_bin", "Z:/not-installed/soffice.exe")

    response = client.post(
        f"/api/resume-versions/{version_id}/exports",
        json={"templateId": TEMPLATE_ID, "format": "pdf"},
    )

    assert response.status_code == 503
    assert response.json()["error"]["code"] == "EXPORT_DEPENDENCY_MISSING"

    with session_factory() as session:
        records = list(session.execute(select(ExportRecord)).scalars().all())
    assert len(records) == 1
    assert records[0].status == "failed"
    assert records[0].error_code == "EXPORT_DEPENDENCY_MISSING"


def test_incomplete_export_cannot_be_downloaded(
    client: TestClient, session_factory, resume_factory
) -> None:
    _, version_id = resume_factory()
    record = ExportRecord(
        id="export_pending",
        resume_version_id=version_id,
        template_id=TEMPLATE_ID,
        format="docx",
        status="processing",
    )
    with session_factory() as session:
        session.add(record)
        session.commit()

    response = client.get("/api/exports/export_pending/download")

    assert response.status_code == 500
    assert response.json()["error"]["code"] == "EXPORT_FAILED"


@pytest.mark.requires_libreoffice
def test_a06_export_pdf_when_libreoffice_available(client: TestClient, resume_factory) -> None:
    if not libreoffice.is_available():
        pytest.skip("本机未安装 LibreOffice，跳过 PDF 导出用例")

    _, version_id = resume_factory()
    response = client.post(
        f"/api/resume-versions/{version_id}/exports",
        json={"templateId": TEMPLATE_ID, "format": "pdf"},
    )

    assert response.status_code == 201, response.text
    data = response.json()["data"]
    assert data["status"] == "completed"

    download = client.get(data["downloadUrl"])
    assert download.status_code == 200
    assert download.content.startswith(b"%PDF")
    assert len(download.content) > 1000
