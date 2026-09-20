"""测试样例：脱敏简历（保持参考简历的排版结构）与文件构造函数。

真实简历与真实 JD 只保存在 .gitignore 覆盖的本地目录，不进入仓库；
这里的姓名、电话、邮箱均为虚构值。
"""

from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document

#: 与参考简历同构的脱敏文本：表格化条目 + 「类别：内容」技能行 + 三行式条目
SAMPLE_TEXT = """姓名：李小明  手机：13800000000  邮箱：lixiaoming@example.com
所在地：杭州  求职意向：后端开发工程师

教育背景
浙江大学 | 计算机科学与技术 | 本科 | 2019.09 - 2023.06

实习经历
杭州云帆科技有限公司 | 后端开发实习生 | 2024.07 - 2024.12
参与订单系统重构，接口平均响应时间下降 40%；
使用 FastAPI 完成后端接口与数据模型开发；

项目经历
ResumeMatch | 后端开发 | 2026.09 - 2026.12
项目背景：面向求职场景的简历匹配系统。
项目内容：负责简历解析、版本管理和导出服务。
成果：完成简历构建与 JD 匹配两条业务闭环。

技能与工具
编程：Python，Java，Git；
AI工具：Trae，VSCode；
协作工具：腾讯文档。
"""

#: 三行式写法：名称行 / 角色行 / 纯日期行
THREE_LINE_STYLE_TEXT = """姓名：王小花
手机：13900000000
教育背景
上海交通大学
软件工程 | 硕士
2023.09 - 至今
实习经历
某互联网公司
后端开发实习生
2024.07 - 2024.12
参与推荐服务开发；
"""

TEXT_PDF_LINES = tuple(
    line for line in SAMPLE_TEXT.splitlines() if line.strip()
)

INFO_TABLE_ROWS = (
    ("姓名：李小明", "手机：13800000000", "邮箱：lixiaoming@example.com"),
    ("所在地：杭州", "求职意向：后端开发工程师", ""),
)

EDUCATION_ROW = ("浙江大学", "计算机科学与技术 | 本科", "2019.09 - 2023.06")

EXPERIENCE_ROWS = (
    ("杭州云帆科技有限公司", "后端开发实习生", "2024.07 - 2024.12"),
    (
        "参与订单系统重构，接口平均响应时间下降 40%；\n使用 FastAPI 完成后端接口与数据模型开发；",
        "",
        "",
    ),
)

PROJECT_ROWS = (
    ("ResumeMatch", "后端开发", "2026.09 - 2026.12"),
    (
        "项目背景：面向求职场景的简历匹配系统。\n项目内容：负责简历解析、版本管理和导出服务。\n成果：完成简历构建与 JD 匹配两条业务闭环。",
        "",
        "",
    ),
)

SKILL_LINES = ("编程：Python，Java，Git；", "AI工具：Trae，VSCode；", "协作工具：腾讯文档。")


def sample_document_dict() -> dict:
    """手工录入场景下的一份合法 ResumeDocument。"""

    return {
        "basicInfo": {
            "name": "李小明",
            "phone": "13800000000",
            "email": "lixiaoming@example.com",
            "location": "杭州",
            "jobTarget": "后端开发工程师",
        },
        "sections": [
            {
                "id": "section_education",
                "type": "education",
                "title": "教育背景",
                "order": 0,
                "items": [
                    {
                        "id": "item_education_1",
                        "order": 0,
                        "fields": {
                            "school": "浙江大学",
                            "major": "计算机科学与技术",
                            "degree": "本科",
                            "startDate": "2019-09",
                            "endDate": "2023-06",
                        },
                        "content": "",
                    }
                ],
            },
            {
                "id": "section_experience",
                "type": "experience",
                "title": "实习经历",
                "order": 1,
                "items": [
                    {
                        "id": "item_experience_1",
                        "order": 0,
                        "fields": {
                            "company": "杭州云帆科技有限公司",
                            "role": "后端开发实习生",
                            "startDate": "2024-07",
                            "endDate": "2024-12",
                        },
                        "content": "参与订单系统重构，接口平均响应时间下降 40%；",
                    }
                ],
            },
            {
                "id": "section_project",
                "type": "project",
                "title": "项目经历",
                "order": 2,
                "items": [
                    {
                        "id": "item_project_1",
                        "order": 0,
                        "fields": {
                            "name": "ResumeMatch",
                            "role": "后端开发",
                            "startDate": "2026-09",
                            "endDate": "2026-12",
                        },
                        "content": "项目背景：面向求职场景的简历匹配系统。\n项目内容：负责简历解析、版本管理和导出服务。",
                    }
                ],
            },
            {
                "id": "section_skill",
                "type": "skill",
                "title": "技能与工具",
                "order": 3,
                "items": [
                    {
                        "id": "item_skill_1",
                        "order": 0,
                        "fields": {"category": "编程"},
                        "content": "Python，Java，Git",
                    },
                    {
                        "id": "item_skill_2",
                        "order": 1,
                        "fields": {"category": "AI工具"},
                        "content": "Trae，VSCode",
                    },
                ],
            },
        ],
    }


def _fill_table(table, rows) -> None:
    for row_index, values in enumerate(rows):
        for column_index, text in enumerate(values):
            if text:
                table.cell(row_index, column_index).text = text


def write_reference_style_docx(path: Path) -> Path:
    """构造与参考简历同构的 DOCX：基本信息表格 + 条目表格 + 技能行。"""

    document = Document()

    _fill_table(document.add_table(rows=len(INFO_TABLE_ROWS), cols=3), INFO_TABLE_ROWS)

    document.add_paragraph("教育背景")
    _fill_table(document.add_table(rows=1, cols=3), (EDUCATION_ROW,))

    document.add_paragraph("实习经历")
    _fill_table(document.add_table(rows=len(EXPERIENCE_ROWS), cols=3), EXPERIENCE_ROWS)

    document.add_paragraph("项目经历")
    _fill_table(document.add_table(rows=len(PROJECT_ROWS), cols=3), PROJECT_ROWS)

    document.add_paragraph("技能与工具")
    for line in SKILL_LINES:
        document.add_paragraph(line)

    document.save(str(path))
    return path


def write_text_pdf(path: Path, lines: tuple[str, ...] = TEXT_PDF_LINES) -> Path:
    document = pymupdf.open()
    page = document.new_page()
    y = 60.0
    for line in lines:
        # 内置中文字体（china-s）保证中文可被 get_text() 提取
        page.insert_text((48, y), line, fontname="china-s", fontsize=10)
        y += 16
    document.save(str(path))
    document.close()
    return path


def write_scanned_pdf(path: Path) -> Path:
    """没有可提取文本的 PDF（模拟扫描件）。"""

    document = pymupdf.open()
    page = document.new_page()
    page.draw_rect(
        pymupdf.Rect(60, 60, 500, 700), color=(0.5, 0.5, 0.5), fill=(0.9, 0.9, 0.9)
    )
    document.save(str(path))
    document.close()
    return path


def long_resume_text(repeat: int = 120) -> str:
    """长文本样例：多模块、多条目，用于验证 order 连续性与性能余量。"""

    lines = [
        "姓名：张三",
        "手机：13700000000",
        "教育背景",
        "某某大学 | 软件工程 | 硕士 | 2019.09 - 2022.06",
        "实习经历",
    ]
    for index in range(repeat):
        lines.append(f"某科技有限公司{index} | 后端开发实习生 | 2024.0{(index % 9) + 1} - 2024.12")
        lines.append("参与服务端接口开发，覆盖日均十万级调用；")
    lines.append("项目经历")
    for index in range(repeat):
        lines.append(f"分布式任务调度平台{index} | 负责人 | 2023.01 - 2023.12")
        lines.append("项目内容：使用 Python 与消息队列完成任务编排。")
    lines.append("技能与工具")
    lines.append("编程：Python，Go，SQL，Redis，Kafka，Docker，Kubernetes；")
    return "\n".join(lines)


def find_reference_resume(repo_root: Path) -> Path | None:
    """查找本机参考简历原件（不入库）；找不到时相关用例跳过。"""

    for candidate in sorted(repo_root.glob("*.docx")):
        if not candidate.name.startswith("~$"):
            return candidate
    return None
